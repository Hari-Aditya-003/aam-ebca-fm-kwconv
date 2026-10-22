#!/usr/bin/env python3
"""Build the project technical DOCX from its version-controlled Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
MARKDOWN = PROJECT_ROOT / "Documentation" / "AAM_EBCA_FM_KWConv_Technical_Documentation.md"
OUTPUT = PROJECT_ROOT / "Documentation" / "AAM_EBCA_FM_KWConv_Technical_Documentation.docx"


def set_cell_shading(cell, colour: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), colour)
    properties.append(shading)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def plain(text: str) -> str:
    text = text.replace("**", "").replace("`", "")
    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)


def add_formula(document: Document, source: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(source.strip())
    run.font.name = "Cambria Math"
    run.font.size = Pt(11)
    run.italic = True


def add_markdown_table(document: Document, rows: list[str]) -> None:
    parsed = [[plain(cell.strip()) for cell in row.strip().strip("|").split("|")] for row in rows]
    if len(parsed) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in parsed[1]):
        parsed.pop(1)
    columns = max(len(row) for row in parsed)
    table = document.add_table(rows=len(parsed), cols=columns)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(parsed):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = values[column_index] if column_index < len(values) else ""
            cell.text = text
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if row_index == 0:
                set_cell_shading(cell, "176B62")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "EAF4F2")
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    colours = ("176B62", "22577A", "2F4858")
    for index, colour in enumerate(colours, 1):
        style = styles[f"Heading {index}"]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor.from_string(colour)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if index == 1 else 8)
        style.paragraph_format.space_after = Pt(5)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 3"].font.size = Pt(11)

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font.size = Pt(8.5)
        code.paragraph_format.left_indent = Inches(0.25)
        code.paragraph_format.right_indent = Inches(0.25)
        code.paragraph_format.space_after = Pt(1)


def build() -> Path:
    if not MARKDOWN.is_file():
        raise FileNotFoundError(MARKDOWN)
    lines = MARKDOWN.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    configure_styles(document)

    header = section.header.paragraphs[0]
    header.text = "AAM–EBCA–FM–KWConv | Technical Documentation"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 100, 110)
    add_page_number(section.footer.paragraphs[0])

    index = 0
    title_written = False
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            paragraph = document.add_paragraph(style="Code Block")
            paragraph.add_run("\n".join(code_lines))
            set_cell = OxmlElement("w:shd")
            set_cell.set(qn("w:fill"), "F3F5F6")
            paragraph._p.get_or_add_pPr().append(set_cell)
            index += 1
            continue

        if stripped == "\\[":
            index += 1
            formula = []
            while index < len(lines) and lines[index].strip() != "\\]":
                formula.append(lines[index].strip())
                index += 1
            add_formula(document, " ".join(formula))
            index += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(document, table_lines)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = plain(heading.group(2))
            if level == 1 and not title_written:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(70)
                run = paragraph.add_run(text)
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(24)
                run.font.color.rgb = RGBColor.from_string("176B62")
                title_written = True
            else:
                document.add_heading(text, level=min(level, 3))
            index += 1
            continue

        if stripped.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.right_indent = Inches(0.25)
            run = paragraph.add_run(plain(stripped[2:]))
            run.italic = True
            run.font.color.rgb = RGBColor.from_string("7A3E00")
            index += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            document.add_paragraph(plain(bullet.group(1)), style="List Bullet")
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            document.add_paragraph(plain(numbered.group(1)), style="List Number")
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", "|", "- ", "> ", "```", "\\[")) or re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        document.add_paragraph(plain(" ".join(paragraph_lines)))

    core = document.core_properties
    core.title = "AAM–EBCA–FM–KWConv Green-Apple Detection System"
    core.subject = "Formula, architecture, code connection, training, depth, and tracking documentation"
    core.author = "Project technical documentation"
    core.keywords = "YOLOv8, AAM, EBCA, focal modulation, KWConv, AdaBins, ByteTrack, green apple"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
